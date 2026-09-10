#!/usr/bin/env python3
"""Build the RMIT Society architecture report as an editable DOCX without Pandoc."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

ROOT = Path(__file__).parent
SOURCE = ROOT / "main.tex"
BIBLIOGRAPHY = ROOT / "references.bib"
OUTPUT = ROOT / "main.docx"
FIGURE_WIDTH = Inches(6.3)
RED = RGBColor(161, 45, 45)
FIGURE_REFS = {
    "fig:roles": "1",
    "fig:overview": "2",
    "fig:core-flow": "3",
    "fig:erd": "4",
    "fig:media-location-flow": "A.1",
    "fig:moderation-flow": "A.2",
    "fig:analytics-flow": "A.3",
    "app:detailed-sequences": "A",
}
TABLE_REFS = {
    "tab:related-work": "1",
    "tab:operation-coverage": "2",
    "tab:components": "3",
}


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_margins(cell, top: int = 70, start: int = 70, bottom: int = 70, end: int = 70) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def add_page_number(section, number_format: str, start: int | None) -> None:
    properties = section._sectPr
    number_type = properties.find(qn("w:pgNumType"))
    if number_type is None:
        number_type = OxmlElement("w:pgNumType")
        properties.append(number_type)
    number_type.set(qn("w:fmt"), number_format)
    if start is not None:
        number_type.set(qn("w:start"), str(start))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def add_toc(paragraph) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'TOC \\o "1-2" \\h \\z \\u')
    paragraph._p.append(field)


def set_update_fields_on_open(document: Document) -> None:
    settings = document.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    content = OxmlElement("w:t")
    content.text = text
    run.append(content)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def extract_braced(text: str, opening: int) -> tuple[str, int]:
    """Return a balanced braced value beginning at ``opening`` and its end index."""
    if text[opening] != "{":
        raise ValueError(f"Expected '{{' at {opening}: {text!r}")
    depth = 0
    for index in range(opening, len(text)):
        if text[index] == "{":
            depth += 1
        elif text[index] == "}":
            depth -= 1
            if depth == 0:
                return text[opening + 1 : index], index + 1
    raise ValueError(f"Unclosed brace in: {text!r}")


def latex_plain(text: str) -> str:
    text = text.replace(r"\liveurl", "<browser-verified live URL>")
    text = text.replace(r"\snapshotdate", "<architecture snapshot date>")
    text = text.replace("~", " ")
    text = re.sub(r"\\ref\*?\{([^}]+)\}", lambda match: FIGURE_REFS.get(match.group(1), TABLE_REFS.get(match.group(1), match.group(1))), text)
    text = re.sub(r"\\shortstack\[[^]]*\]\{(.*?)\}", lambda match: match.group(1).replace(r"\\", " "), text)
    text = text.replace(r"\\", " ")
    text = text.replace(r"\_", "_").replace(r"\&", "&").replace(r"\%", "%")
    text = text.replace(r"\#", "#").replace(r"\$", "$")
    text = text.replace("$", "")
    return text


def add_inline(paragraph, text: str, citations: dict[str, int], *, italic: bool = False, bold: bool = False, color=None) -> None:
    command = re.compile(r"\\(texttt|emph|textbf|url|cite|hyperref|shortstack)(?:\[[^]]*\])?\{")
    cursor = 0
    while match := command.search(text, cursor):
        before = latex_plain(text[cursor : match.start()])
        if before:
            run = paragraph.add_run(before)
            run.italic = italic
            run.bold = bold
            if color:
                run.font.color.rgb = color
        value, end = extract_braced(text, match.end() - 1)
        kind = match.group(1)
        if kind == "cite":
            labels = [label.strip() for label in value.split(",")]
            run = paragraph.add_run("[" + ", ".join(str(citations[label]) for label in labels) + "]")
            run.italic = italic
            run.bold = bold
        elif kind == "url":
            add_hyperlink(paragraph, latex_plain(value), latex_plain(value))
        elif kind == "texttt":
            run = paragraph.add_run(latex_plain(value))
            run.font.name = "Courier New"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
            run.italic = italic
            run.bold = bold
        elif kind == "emph":
            add_inline(paragraph, value, citations, italic=True, bold=bold, color=color)
        elif kind == "textbf":
            add_inline(paragraph, value, citations, italic=italic, bold=True, color=color)
        elif kind == "shortstack":
            add_inline(paragraph, value.replace(r"\\", " "), citations, italic=italic, bold=bold, color=color)
        else:
            add_inline(paragraph, value, citations, italic=italic, bold=bold, color=color)
        cursor = end
    after = latex_plain(text[cursor:])
    if after:
        run = paragraph.add_run(after)
        run.italic = italic
        run.bold = bold
        if color:
            run.font.color.rgb = color


def add_body_paragraph(document: Document, text: str, citations: dict[str, int], *, style: str = "Normal", italic: bool = False) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.08
    add_inline(paragraph, text.strip(), citations, italic=italic)


def clean_table_line(line: str) -> str:
    line = line.strip()
    for token in (r"\toprule", r"\midrule", r"\bottomrule", r"\endfirsthead", r"\endhead"):
        line = line.replace(token, "")
    return line.strip()

def split_latex_rows(line: str) -> list[str]:
    rows: list[str] = []
    start = 0
    depth = 0
    index = 0
    while index < len(line):
        if line[index] == "{":
            depth += 1
        elif line[index] == "}":
            depth -= 1
        elif line.startswith(r"\\", index) and depth == 0:
            rows.append(line[start:index])
            index += 2
            start = index
            continue
        index += 1
    rows.append(line[start:])
    return rows


def add_table(document: Document, lines: list[str], citations: dict[str, int]) -> None:
    caption = ""
    rows: list[list[str]] = []
    in_repeated_header = False
    for raw_line in lines:
        if r"\endfirsthead" in raw_line:
            in_repeated_header = True
            continue
        if in_repeated_header:
            if r"\endhead" in raw_line:
                in_repeated_header = False
            continue
        if r"\caption{" in raw_line:
            opening = raw_line.index("{", raw_line.index(r"\caption"))
            caption, _ = extract_braced(raw_line, opening)
        line = clean_table_line(raw_line)
        if not line or line.startswith(r"\caption"):
            continue
        for segment in split_latex_rows(line):
            segment = clean_table_line(segment)
            if not re.search(r"(?<!\\)&", segment):
                continue
            cells = [cell.strip() for cell in re.split(r"(?<!\\)&", segment)]
            rows.append(cells)
    if caption:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(latex_plain(caption))
        run.italic = True
        run.font.size = Pt(9)
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=1, cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for row_index, values in enumerate(rows):
        cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
        if row_index == 0:
            set_repeat_table_header(table.rows[0])
        for column_index, cell in enumerate(cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1
            add_inline(paragraph, values[column_index] if column_index < len(values) else "", citations)
            for run in paragraph.runs:
                run.font.size = Pt(8.5 if column_count >= 3 else 9)
                if "<browser-verified live URL>" in run.text:
                    run.font.name = "Courier New"
                    run.font.color.rgb = RED
            if row_index == 0:
                set_cell_shading(cell, "E7E6E6")
                for run in paragraph.runs:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_bibliography() -> dict[str, dict[str, str]]:
    entries: dict[str, dict[str, str]] = {}
    content = BIBLIOGRAPHY.read_text(encoding="utf-8")
    for match in re.finditer(r"@\w+\{([^,]+),\s*(.*?)\}\s*(?=@|\Z)", content, re.DOTALL):
        key, body = match.groups()
        fields = {
            field: latex_plain(value)
            for field, value in re.findall(r"(\w+)\s*=\s*\{(.*?)\}(?:,|$)", body, re.DOTALL)
        }
        entries[key] = fields
    return entries


def citations_in_order(body: str) -> dict[str, int]:
    citations: dict[str, int] = {}
    for match in re.finditer(r"\\cite\{([^}]+)\}", body):
        for key in (part.strip() for part in match.group(1).split(",")):
            citations.setdefault(key, len(citations) + 1)
    return citations


def add_cover(document: Document) -> None:
    def rule(width: int) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(18)
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), str(width))
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "000000")
        border.append(bottom)
        paragraph._p.get_or_add_pPr().append(border)

    rule(18)
    university = document.add_paragraph()
    university.paragraph_format.space_after = Pt(10)
    run = university.add_run("RMIT UNIVERSITY")
    run.bold = True
    run.font.size = Pt(26)
    school = document.add_paragraph("School of Computing Technologies")
    school.runs[0].font.size = Pt(16)
    document.add_paragraph().paragraph_format.space_after = Pt(120)
    rule(7)
    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(20)
    run = title.add_run("RMIT Society\nCloud Solution Architecture")
    run.bold = True
    run.font.size = Pt(26)
    assessment = document.add_paragraph("Assessment 3 — Cloud Computing")
    assessment.runs[0].font.size = Pt(16)
    document.add_paragraph().paragraph_format.space_after = Pt(35)
    details = [
        ("Course code and name:", "<course code and course name>"),
        ("Student name:", "<student name>"),
        ("Student ID:", "<student ID>"),
        ("Tutor:", "<tutor name>"),
        ("Submission date:", "<submission date>"),
        ("Report version:", "0.1 — architecture-report draft"),
        ("Architecture snapshot:", "<architecture snapshot date>"),
    ]
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for label, value in details:
        cells = table.add_row().cells
        cells[0].width = Inches(1.65)
        cells[1].width = Inches(4.55)
        label_run = cells[0].paragraphs[0].add_run(label)
        label_run.bold = True
        value_run = cells[1].paragraphs[0].add_run(value)
        if value.startswith("<"):
            value_run.font.name = "Courier New"
            value_run.font.color.rgb = RED
        for cell in cells:
            set_cell_margins(cell, 30, 30, 30, 30)
    document.add_paragraph().paragraph_format.space_after = Pt(70)
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(18)
    note.paragraph_format.space_after = Pt(0)
    note.paragraph_format.left_indent = Inches(0.2)
    note.paragraph_format.right_indent = Inches(0.2)
    run = note.add_run(
        "This boilerplate cover intentionally contains editable placeholders. Replace the student "
        "and assessment details, then remove placeholder styling before submission."
    )
    run.italic = True
    run.font.size = Pt(9)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(24)
    section.bottom_margin = Mm(24)
    section.left_margin = Mm(24)
    section.right_margin = Mm(24)
    section.different_first_page_header_footer = True
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    for name, size in (("Title", 26), ("Heading 1", 16), ("Heading 2", 13)):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
    if "Figure Caption" not in document.styles:
        style = document.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Normal"]
        style.font.size = Pt(9)
        style.font.italic = True
        style.paragraph_format.space_after = Pt(8)
    set_update_fields_on_open(document)


def add_figure(document: Document, lines: list[str], citations: dict[str, int]) -> None:
    source = next((re.search(r"\{(figures/[^}]+)\}", line).group(1) for line in lines if r"\includegraphics" in line), None)
    caption = next(
        (
            extract_braced(line, line.index("{", line.index(r"\caption")))[0]
            for line in lines
            if r"\caption{" in line
        ),
        None,
    )
    if source:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(ROOT / source), width=FIGURE_WIDTH)
    if caption:
        paragraph = document.add_paragraph(style="Figure Caption")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(paragraph, caption, citations)

def add_references(document: Document, citations: dict[str, int]) -> None:
    bibliography = parse_bibliography()
    document.add_page_break()
    document.add_heading("References", level=1)
    for key, number in citations.items():
        entry = bibliography[key]
        author = entry.get("author", "")
        author = author.strip("{}")
        title = entry.get("title", "")
        url_match = re.search(r"\\url\{([^}]+)\}", entry.get("howpublished", ""))
        url = url_match.group(1) if url_match else entry.get("howpublished", "")
        note = entry.get("note", "")
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.add_run(f"[{number}] {author}, “{title},” {entry.get('year', '')}. ")
        if url:
            add_hyperlink(paragraph, url, url)
        if note:
            paragraph.add_run(f" ({note}).")


def build() -> None:
    source = SOURCE.read_text(encoding="utf-8")
    body = source.split(r"\pagenumbering{arabic}", 1)[1].split(r"\bibliographystyle", 1)[0]
    citations = citations_in_order(body)
    document = Document()
    configure_document(document)
    add_cover(document)

    document.add_page_break()
    toc_section = document.add_section(WD_SECTION.CONTINUOUS)
    add_page_number(toc_section, "lowerRoman", 1)
    document.add_heading("Table of Contents", level=1)
    toc = document.add_paragraph()
    add_toc(toc)

    document.add_page_break()
    report_section = document.add_section(WD_SECTION.CONTINUOUS)
    add_page_number(report_section, "decimal", 1)
    lines = body.splitlines()
    index = 0
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_lines
        if paragraph_lines:
            text = " ".join(line.strip() for line in paragraph_lines)
            add_body_paragraph(document, text, citations)
            paragraph_lines = []

    while index < len(lines):
        line = lines[index].strip()
        if not line:
            flush_paragraph()
            index += 1
            continue
        if line.startswith(r"\section{"):
            flush_paragraph()
            title, _ = extract_braced(line, line.index("{"))
            document.add_heading(latex_plain(title), level=1)
        elif line.startswith(r"\subsection{"):
            flush_paragraph()
            title, _ = extract_braced(line, line.index("{"))
            document.add_heading(latex_plain(title), level=2)
        elif line.startswith(r"\begin{figure}"):
            flush_paragraph()
            end = next(position for position in range(index, len(lines)) if lines[position].strip() == r"\end{figure}")
            add_figure(document, lines[index : end + 1], citations)
            document.add_page_break()
            index = end
        elif line.startswith(r"\begin{table}") or line.startswith(r"\begin{tabularx}") or line.startswith(r"\begin{longtable}"):
            flush_paragraph()
            environment = "longtable" if line.startswith(r"\begin{longtable}") else "tabularx"
            end_marker = rf"\end{{{environment}}}"
            end = next(position for position in range(index, len(lines)) if lines[position].strip() == end_marker)
            add_table(document, lines[index : end + 1], citations)
            if line.startswith(r"\begin{table}"):
                index = next(position for position in range(end, len(lines)) if lines[position].strip() == r"\end{table}")
            else:
                index = end
        elif line == r"\clearpage":
            flush_paragraph()
            document.add_page_break()
        elif line == r"\appendix":
            flush_paragraph()
            document.add_page_break()
        elif line == r"\statusnote":
            flush_paragraph()
            add_body_paragraph(
                document,
                "Deployment-status note: replace every provisional statement in this report with evidence from the final demonstration environment before submission.",
                citations,
                italic=True,
            )
        elif line.startswith((r"\label", r"\setcounter", r"\renewcommand", r"\small", r"\normalsize")):
            flush_paragraph()
        else:
            paragraph_lines.append(line)
        index += 1

    flush_paragraph()
    add_references(document, citations)
    document.save(OUTPUT)


if __name__ == "__main__":
    build()
